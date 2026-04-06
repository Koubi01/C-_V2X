#!/usr/bin/env python3
"""Generate a static SV-by-road-reference dataset for the MapLibre traffic layer.

This script reads the Czech workbook V2_CSD_2025.xlsx and creates a compact JSON
artifact used by the client map overlay.

Layer design choices:
- Metric: SV (all motor vehicles/day)
- Geometry source: existing OpenMapTiles `transportation` layer
- Matching key: road reference (`SIL` -> OSM `ref`)
- Confidence policy: all road refs with SV values are emitted by default
"""

from __future__ import annotations

import argparse
import json
import math
import statistics
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class BinSpec:
    id: str
    min_inclusive: int
    max_exclusive: int | None
    label: str
    color: str


BINS: list[BinSpec] = [
    BinSpec("traffic_sv_1", 0, 2001, "SV <= 2,000", "#fde68a"),
    BinSpec("traffic_sv_2", 2001, 5001, "2,001 - 5,000", "#fbbf24"),
    BinSpec("traffic_sv_3", 5001, 10001, "5,001 - 10,000", "#fb923c"),
    BinSpec("traffic_sv_4", 10001, 20001, "10,001 - 20,000", "#f97316"),
    BinSpec("traffic_sv_5", 20001, None, "SV > 20,000", "#b91c1c"),
]

EXCLUDED_REFS = {"", "-"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate static traffic SV layer dataset.")
    parser.add_argument(
        "--input",
        default="V2_CSD_2025.xlsx",
        help="Path to the source workbook (default: V2_CSD_2025.xlsx)",
    )
    parser.add_argument(
        "--output",
        default="src/V2XDashboard.Client/wwwroot/data/traffic-sv-by-roadref.json",
        help="Output JSON path for client layer data",
    )
    parser.add_argument(
        "--min-samples",
        type=int,
        default=1,
        help="Minimum sample count per road ref to consider it confident",
    )
    parser.add_argument(
        "--max-cv",
        type=float,
        default=999.0,
        help="Maximum coefficient of variation (stdev/mean) for confidence",
    )
    return parser.parse_args()


def normalize_ref(value: object) -> str:
    if value is None:
        return ""

    normalized = str(value).strip().upper()
    if normalized.isdigit():
        return str(int(normalized))

    return normalized


def parse_sv(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, str):
        candidate = value.strip().replace(" ", "")
        if not candidate or candidate == "-":
            return None
        try:
            return int(float(candidate.replace(",", ".")))
        except ValueError:
            return None

    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return int(value)

    return None


def pick_bin(value: int) -> BinSpec:
    for bin_spec in BINS:
        if bin_spec.max_exclusive is None:
            if value >= bin_spec.min_inclusive:
                return bin_spec
        elif bin_spec.min_inclusive <= value < bin_spec.max_exclusive:
            return bin_spec

    return BINS[-1]


def build_header_map(header_row: Iterable[object]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for idx, value in enumerate(header_row, start=1):
        if value is None:
            continue
        key = str(value).strip()
        if key:
            header_map[key] = idx
    return header_map


def main() -> int:
    args = parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")

    wb = load_workbook(input_path, data_only=True)
    if "V2_CSD2025" not in wb.sheetnames:
        raise ValueError("Expected sheet 'V2_CSD2025' not found in workbook.")

    ws = wb["V2_CSD2025"]
    headers = [ws.cell(2, col).value for col in range(1, ws.max_column + 1)]
    header_map = build_header_map(headers)

    required_columns = ["SIL", "SV"]
    missing = [name for name in required_columns if name not in header_map]
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(missing)}")

    sil_col = header_map["SIL"]
    sv_col = header_map["SV"]

    grouped_sv: dict[str, list[int]] = {}

    total_rows = 0
    rows_with_sv = 0
    for row in range(3, ws.max_row + 1):
        total_rows += 1

        road_ref = normalize_ref(ws.cell(row, sil_col).value)
        if road_ref in EXCLUDED_REFS:
            continue

        sv_value = parse_sv(ws.cell(row, sv_col).value)
        if sv_value is None:
            continue

        rows_with_sv += 1
        grouped_sv.setdefault(road_ref, []).append(sv_value)

    roads: list[dict[str, object]] = []
    hidden_low_confidence = 0

    for road_ref, values in grouped_sv.items():
        if len(values) < args.min_samples:
            hidden_low_confidence += 1
            continue

        mean_sv = float(statistics.fmean(values))
        stdev_sv = float(statistics.pstdev(values)) if len(values) > 1 else 0.0
        cv = stdev_sv / mean_sv if mean_sv > 0 else 0.0
        if cv > args.max_cv:
            hidden_low_confidence += 1
            continue

        median_sv = int(round(statistics.median(values)))
        bin_spec = pick_bin(median_sv)

        roads.append(
            {
                "ref": road_ref,
                "svMedian": median_sv,
                "svMean": int(round(mean_sv)),
                "sampleCount": len(values),
                "coefficientOfVariation": round(cv, 4),
                "classId": bin_spec.id,
                "label": bin_spec.label,
                "color": bin_spec.color,
            }
        )

    roads.sort(key=lambda item: str(item["ref"]))

    payload = {
        "generatedAtUtc": datetime.now(timezone.utc).isoformat(),
        "sourceFile": input_path.name,
        "metric": "SV",
        "matchKey": "SIL->OSM:ref",
        "legend": [
            {
                "id": item.id,
                "label": item.label,
                "minInclusive": item.min_inclusive,
                "maxExclusive": item.max_exclusive,
                "color": item.color,
            }
            for item in BINS
        ],
        "roads": roads,
        "metadata": {
            "totalRows": total_rows,
            "rowsWithSv": rows_with_sv,
            "distinctRoadRefs": len(grouped_sv),
            "emittedRoadRefs": len(roads),
            "hiddenLowConfidenceRoadRefs": hidden_low_confidence,
            "confidencePolicy": {
                "minSamples": args.min_samples,
                "maxCoefficientOfVariation": args.max_cv,
            },
        },
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Generated {len(roads)} road refs to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
